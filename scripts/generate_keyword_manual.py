"""
generate_keyword_manual.py - Generates Rakshak Node Keyword Manual as a .docx file
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT_PATH = r"C:\Users\Hp\Desktop\Rakshak_Node_Keyword_Manual.docx"

def set_cell_background(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_border(cell, color="D1D5DB"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tcBorders.append(border)
    tcPr.append(tcBorders)

doc = Document()

# Page margins
section = doc.sections[0]
section.top_margin    = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin   = Cm(2.5)
section.right_margin  = Cm(2.5)

# ── Title ──────────────────────────────────────────────────────────
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("🛡️  RAKSHAK")
run.font.size = Pt(28)
run.font.bold = True
run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

doc.add_paragraph()

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = subtitle.add_run("ON-CHAIN IMMUNITY SYSTEM")
sub.font.size = Pt(13)
sub.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
sub.font.bold = True

doc.add_paragraph()

title2 = doc.add_paragraph()
title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
t2 = title2.add_run("Node Keyword Reference Manual")
t2.font.size = Pt(18)
t2.font.bold = True
t2.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

doc.add_paragraph()

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
i = intro.add_run(
    "Every node in the Rakshak 3D Galaxy is classified with a behavioral archetype keyword.\n"
    "This manual explains what each keyword means in plain English."
)
i.font.size = Pt(10)
i.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

doc.add_paragraph()

# ── Section helper ─────────────────────────────────────────────────
def add_section_heading(doc, title, color_hex):
    h = doc.add_paragraph()
    run = h.add_run(f"  {title}")
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(2)
    # Shade the paragraph background via XML
    pPr = h._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    r, g, b = int(color_hex[0:2],16), int(color_hex[2:4],16), int(color_hex[4:6],16)
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    pPr.append(shd)
    return h

def add_keyword_table(doc, rows):
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # Header
    hdr = table.rows[0].cells
    headers = ['Icon', 'Keyword', 'Plain English Meaning']
    header_colors = ['1E293B', '1E293B', '1E293B']
    for i, (cell, text, clr) in enumerate(zip(hdr, headers, header_colors)):
        set_cell_background(cell, clr)
        set_cell_border(cell, '334155')
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
        run.font.size = Pt(10)

    # Set column widths
    widths = [Cm(1.5), Cm(5.5), Cm(10)]
    for i, width in enumerate(widths):
        for cell in table.columns[i].cells:
            cell.width = width

    # Data rows
    for row_idx, (icon, keyword, meaning, row_color) in enumerate(rows):
        row = table.rows[row_idx + 1]
        data = [icon, keyword, meaning]
        for col_idx, (cell, text) in enumerate(zip(row.cells, data)):
            set_cell_background(cell, row_color)
            set_cell_border(cell, 'CBD5E1')
            p = cell.paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.font.size = Pt(10)
            if col_idx == 1:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            else:
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    doc.add_paragraph()

# ── SECTION 1: THREAT / HIGH RISK ─────────────────────────────────
add_section_heading(doc, "🔴  THREAT / HIGH RISK  (Risk Score > 65%)", "7F1D1D")

threat_rows = [
    ("⚡", "RAPID EXFILTRATOR",
     "Actively draining a victim wallet — rapid extraction and dispersal of funds across many addresses. Classic hack/exploit pattern.",
     "FEF2F2"),
    ("🌪️", "TORNADO MIXER RELAY",
     "Routing money through privacy mixing pools (like Tornado Cash) to hide where it originally came from — money laundering behavior.",
     "FFF7ED"),
    ("🔄", "CIRCULAR WASHTRADER",
     "Sending money in circles (Wallet A → B → C → A) to fake trading volume and manipulate market metrics.",
     "FFF1F2"),
    ("🐋", "WHALE DUMP LIQUIDATOR",
     "A huge wallet suddenly dumping massive amounts of tokens — causes severe price crashes for other holders.",
     "FEF2F2"),
    ("🤖", "MEV SANDWICH BOT",
     "Automated bot that sneaks its own transaction in front of yours (frontrunning) to steal your trade profit using gas bidding.",
     "FDF4FF"),
    ("📦", "SYBIL AIRDROP OPERATOR",
     "Hundreds of fake wallet addresses coordinated by one entity to abuse airdrop and reward systems — scripted bulk farming.",
     "FFF1F2"),
    ("🕸️", "TAINTED GRAPH RELAY",
     "This wallet isn't directly doing something bad, but it's connected to flagged wallets nearby. Suspicious by association.",
     "FFFBEB"),
]
add_keyword_table(doc, threat_rows)

# ── SECTION 2: MEDIUM / WATCH ──────────────────────────────────────
add_section_heading(doc, "🟡  MEDIUM / WATCH  (Risk Score 25–65%)", "78350F")

medium_rows = [
    ("⚡", "FLASH LOAN OPERATOR",
     "Borrows millions of dollars instantly with zero collateral in a single blockchain transaction — often used in DeFi protocol attacks.",
     "FFFBEB"),
    ("🏛️", "INSTITUTIONAL GATEWAY",
     "A high-throughput wallet processing large volumes — likely a centralized exchange (CEX) or institutional clearing hub.",
     "F5F3FF"),
    ("🎨", "NFT HIGH-VELOCITY TRADER",
     "Buying and selling NFTs at very high speed — can indicate wash trading to fake NFT collection floor prices.",
     "FAF5FF"),
]
add_keyword_table(doc, medium_rows)

# ── SECTION 3: SAFE / TRUSTED ──────────────────────────────────────
add_section_heading(doc, "🔵  SAFE / TRUSTED  (Risk Score < 25%)", "0C4A6E")

safe_rows = [
    ("🛡️", "GRADUATED SENTINEL USER",
     "Fully verified clean wallet — 45+ consecutive days of honest on-chain activity. Has passed all Rakshak trust graduation checks.",
     "F0FDF4"),
    ("🌱", "NEW ON-RAMP WALLET",
     "A brand new wallet with little history. Not flagged — being watched carefully under Rakshak's new-wallet protection caps.",
     "ECFDF5"),
    ("🌾", "BLUE-CHIP YIELD FARMER",
     "Actively providing liquidity to audited, trusted DeFi protocols (Uniswap, Aave, Compound). Normal participation.",
     "F0FDFA"),
    ("🌉", "CROSS-CHAIN BRIDGE RELAY",
     "Moving assets between different blockchains (e.g. Ethereum → Base → Polygon). Normal cross-chain infrastructure wallet.",
     "EFF6FF"),
    ("🏦", "COLD STORAGE WHALE VAULT",
     "Large balance wallet that almost never moves — long-term holder, dormant reserve. Low risk despite high value.",
     "F0F9FF"),
    ("⚙️", "DEFI AMM ARBITRAGEUR",
     "Automated trading bot that balances prices across decentralized exchanges. Mostly harmless market-efficiency mechanism.",
     "ECFEFF"),
]
add_keyword_table(doc, safe_rows)

# ── SECTION 4: COLOR LEGEND ────────────────────────────────────────
doc.add_paragraph()
add_section_heading(doc, "🎨  3D Galaxy Color Legend", "1E3A5F")
doc.add_paragraph()

color_table = doc.add_table(rows=6, cols=3)
color_table.alignment = WD_TABLE_ALIGNMENT.CENTER
color_table.style = 'Table Grid'

# Headers
for i, txt in enumerate(['Color / Glow', 'Risk Level', 'Meaning']):
    c = color_table.rows[0].cells[i]
    set_cell_background(c, '1E293B')
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
    r.font.size = Pt(10)

color_rows = [
    ("🔴  Red / Crimson",   "> 85%  CRITICAL",   "Immediate auto-block. Confirmed malicious behavior.", "FEF2F2"),
    ("🟠  Orange",           "65–85%  ELEVATED",  "Human review required. High confidence threat signals.", "FFF7ED"),
    ("🟡  Amber / Yellow",   "25–65%  WATCH",     "Monitor closely. Unusual patterns, not yet confirmed.", "FFFBEB"),
    ("🔵  Blue / Teal",      "< 25%  LOW",        "Normal activity. Trusted or newly on-boarded wallet.", "EFF6FF"),
    ("🟢  Green",            "< 10%  CLEAR",      "Graduated trusted wallet. Passed all Rakshak trust checks.", "F0FDF4"),
]

for row_idx, (clr, risk, meaning, bg) in enumerate(color_rows):
    row = color_table.rows[row_idx + 1]
    for col_idx, (cell, text) in enumerate(zip(row.cells, [clr, risk, meaning])):
        set_cell_background(cell, bg)
        set_cell_border(cell, 'CBD5E1')
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        if col_idx == 1:
            r.font.bold = True

doc.add_paragraph()

# ── Footer note ────────────────────────────────────────────────────
note = doc.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
n = note.add_run(
    "Rakshak On-Chain Immunity System  •  Built for Nexus Hackathon  •  Base Sepolia (Chain ID: 84532)\n"
    "All classifications are deterministic — the same wallet address always resolves to the same archetype keyword."
)
n.font.size = Pt(9)
n.font.italic = True
n.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

# Save
doc.save(OUTPUT_PATH)
import sys
sys.stdout.buffer.write(f"Document saved to: {OUTPUT_PATH}\n".encode('utf-8'))